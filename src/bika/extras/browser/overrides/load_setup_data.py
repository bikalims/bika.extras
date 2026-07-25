# -*- coding: utf-8 -*-

import tempfile
import traceback

import transaction
from bika.lims import PMF
from bika.lims import logger
from bika.lims.interfaces import ISetupDataImporter
from openpyxl import load_workbook
from pkg_resources import resource_filename
from Products.CMFCore.utils import getToolByName
from senaite.core.catalog import CLIENT_CATALOG
from senaite.core.catalog import SETUP_CATALOG
from senaite.core.exportimport.load_setup_data import LoadSetupData as Base
from zope.component import getAdapters
from zope.component.hooks import getSite


class LoadSetupData(Base):
    """Load setup data without rebuilding unrelated catalogs."""

    def __call__(self):
        form = self.request.form
        portal = getSite()
        workbook = None

        if ("setupexisting" in form and "existing" in form
                and form["existing"]):
            fn = form["existing"].split(":")
            self.dataset_project = fn[0]
            self.dataset_name = fn[1]
            path = "setupdata/%s/%s.xlsx" % (
                self.dataset_name, self.dataset_name)
            filename = resource_filename(self.dataset_project, path)
            try:
                workbook = load_workbook(filename=filename)
            except AttributeError:
                print("")
                print(traceback.format_exc())
                print("Error while loading ", path)

        elif ("setupfile" in form and "file" in form and form["file"]
              and "projectname" in form and form["projectname"]):
            self.dataset_project = form["projectname"]
            tmp = tempfile.mktemp(suffix=".xlsx")
            file_content = form["file"].read()
            open(tmp, "wb").write(file_content)
            workbook = load_workbook(filename=tmp)
            self.dataset_name = "uploaded"

        if not workbook:
            message = PMF("File not found...")
            self.context.plone_utils.addPortalMessage(message)
            self.request.RESPONSE.redirect(portal.absolute_url() + "/import")
            return

        adapters = [
            [name, adapter]
            for name, adapter
            in list(getAdapters((self.context, ), ISetupDataImporter))
        ]
        for sheetname in workbook.sheetnames:
            transaction.savepoint()
            ad_name = sheetname.replace(" ", "_")
            if ad_name in [adapter[0] for adapter in adapters]:
                adapter = [
                    adapter[1] for adapter in adapters
                    if adapter[0] == ad_name
                ][0]
                adapter(
                    self, workbook, self.dataset_project, self.dataset_name)
                adapters = [
                    adapter for adapter in adapters
                    if adapter[0] != ad_name
                ]

        for name, adapter in adapters:
            transaction.savepoint()
            adapter(self, workbook, self.dataset_project, self.dataset_name)

        check = len(self.deferred)
        while self.deferred:
            new = self.solve_deferred()
            logger.info("solved %s of %s deferred references" % (
                check - new, check))
            if new == check:
                raise Exception("%s unsolved deferred references: %s" % (
                    len(self.deferred), self.deferred))
            check = new

        for catalog_name in (SETUP_CATALOG, CLIENT_CATALOG):
            logger.info("Rebuilding %s", catalog_name)
            catalog = getToolByName(self.context, catalog_name)
            catalog.clearFindAndRebuild()
