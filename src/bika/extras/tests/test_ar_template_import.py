# -*- coding: utf-8 -*-
import os
import unittest
from openpyxl import Workbook, load_workbook
from bika.extras.browser.overrides import setupdata


class Record(object):
    def __init__(self, **values):
        self.__dict__.update(values)

    def UID(self):
        return self.uid

    def reindexObject(self):
        self.reindexed = True

    def setSampleType(self, value):
        self.sample_type = value

    def setSamplePoint(self, value):
        self.sample_point = value

    def setPartitions(self, value):
        self.partitions = value

    def setServices(self, value):
        self.services = value


class Field(object):
    def __init__(self, name):
        self.name = name

    def set(self, obj, value):
        setattr(obj, self.name, value)


class TestARTemplateImport(unittest.TestCase):
    def setUp(self):
        self.originals = (setupdata.api.create, setupdata.api.get_fields,
                          setupdata.api.get_senaite_setup,
                          setupdata.getToolByName, setupdata.notify)
        self.created = []
        self.events = []
        self.folder = object()
        self.matrix = Record(uid="matrix-uid")
        self.service = Record(uid="service-uid")
        self.importer = setupdata.AR_Templates(None)
        self.importer.context = Record()
        self.importer.lsd = Record(deferred=[])
        setupdata.api.get_senaite_setup = lambda: Record(sampletemplates=self.folder)
        setupdata.api.get_fields = lambda obj: {
            name: Field(name) for name in ("composite", "target_tat", "matrix_reference")}
        setupdata.getToolByName = lambda *args: object()
        setupdata.notify = lambda event: self.events.append(event.object)

        def create(folder, portal_type, **values):
            self.assertIs(folder, self.folder)
            obj = Record(**values)
            self.created.append(obj)
            return obj

        setupdata.api.create = create
        self.importer.get_object = lambda catalog, portal_type, title: (
            self.matrix if portal_type == "MatrixReference" else
            self.service if portal_type == "AnalysisService" else None)

    def tearDown(self):
        (setupdata.api.create, setupdata.api.get_fields,
         setupdata.api.get_senaite_setup, setupdata.getToolByName,
         setupdata.notify) = self.originals

    def workbook(self, friendly_only=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "AR Templates"
        extra_headers = [None, None, None] if friendly_only else [
            "Composite", "matrix_reference", "target_tat"]
        sheet.append(["title", "Client_title", "description",
                      "SampleType_title", "SamplePoint_title"] + extra_headers)
        sheet.append(["Sample Registration Templates"])
        sheet.append(["Title", "Client", "Description", "Sample Type",
                      "Sample Point", "Composite Y/N", "Matrix Reference", "Target TAT"])
        sheet.append(["SM-01", "lab", "Feed", "Ore", "Feed point", "Y", "SM-01", 8])
        sheet.append(["SM-02", "lab", "Discharge", "Slurry", "Discharge", "N", "SM-02", 0])
        return workbook

    def import_workbook(self, workbook):
        self.importer.lsd.context = self.importer.context
        self.importer(self.importer.lsd, workbook, "bika.extras", "test")

    def test_extra_fields_and_zero_target(self):
        self.import_workbook(self.workbook())
        self.assertEqual(len(self.created), 2)
        self.assertTrue(self.created[0].composite)
        self.assertFalse(self.created[1].composite)
        self.assertEqual(self.created[0].target_tat, 8.0)
        self.assertEqual(self.created[1].target_tat, 0.0)
        self.assertIs(self.created[0].matrix_reference, self.matrix)
        self.assertEqual(self.events, self.created)
        self.assertTrue(self.created[0].reindexed)

    def test_row_three_labels_fill_missing_extra_headers(self):
        self.import_workbook(self.workbook(friendly_only=True))
        self.assertTrue(self.created[0].composite)
        self.assertEqual(self.created[0].target_tat, 8.0)
        self.assertIs(self.created[0].matrix_reference, self.matrix)

    def test_missing_partition_column_and_optional_partition_sheet(self):
        workbook = self.workbook()
        sheet = workbook.create_sheet("AR Template Analyses")
        sheet.append(["ARTemplate", "service_uid"])
        sheet.append(["Description"])
        sheet.append(["Example"])
        sheet.append(["SM-01", "Uranium"])
        self.import_workbook(workbook)
        self.assertEqual(self.created[0].services,
                         [{"uid": "service-uid", "part_id": "part-1"}])
        self.assertEqual(self.created[1].services, [])
        self.assertEqual(self.created[0].partitions[0]["part_id"], "part-1")

    def test_missing_matrix_is_deferred(self):
        self.importer.get_object = lambda *args: None
        self.import_workbook(self.workbook())
        self.assertEqual(len(self.importer.lsd.deferred), 2)
        reference = self.importer.lsd.deferred[0]
        self.assertEqual(reference["src_field"], "matrix_reference")
        self.assertEqual(reference["dest_query"]["title"], "SM-01")

    def test_blank_optional_values_are_not_written(self):
        values = self.importer.get_extra_values({"Composite": "", "Target TAT": ""})
        obj = Record()
        self.importer.set_extra_values(obj, values, None)
        self.assertFalse(hasattr(obj, "composite"))
        self.assertFalse(hasattr(obj, "target_tat"))

    def test_invalid_optional_values_are_rejected(self):
        for value in ("-1", "nan", "inf", "invalid"):
            with self.assertRaises(ValueError):
                self.importer.get_extra_values({"Target TAT": value})
        with self.assertRaises(ValueError):
            self.importer.get_extra_values({"Composite": "maybe"})

    def test_lotus_workbook(self):
        path = os.path.abspath(os.path.join(
            os.path.dirname(__file__), "../../../../../bika.mining/src/bika/mining",
            "setupdata/Lotus new setup data.xlsx"))
        if not os.path.isfile(path):
            self.skipTest("Lotus workbook is available in the mining checkout")
        workbook = load_workbook(path, data_only=True)
        self.importer.worksheet = workbook["AR Templates"]
        rows = [row for row in self.importer.get_rows(3) if row.get("title")]
        self.assertEqual(len(rows), 77)
        setupdata.getToolByName = lambda *args: (
            lambda **query: [Record(getObject=lambda: self.folder)])
        self.import_workbook(workbook)
        self.assertEqual(len(self.created), 77)
        for obj, row in zip(self.created, rows):
            expected = self.importer.get_extra_values(row)
            self.assertEqual(obj.title, row["title"])
            if "target_tat" in expected:
                self.assertEqual(obj.target_tat, expected["target_tat"])
            if expected["matrix_reference"]:
                self.assertIs(obj.matrix_reference, self.matrix)
            if "composite" in expected:
                self.assertEqual(obj.composite, expected["composite"])
